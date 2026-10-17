"""Turning real evidence files into timeline events.

`text_lines(f)` gives every readable file the same shape, a list of lines:
.txt/.csv as written, a chat export inside a .zip, PDF statement text, Excel
rows, and OCR text from images (Tesseract, English). Every event cites
{file, line} in that list, and the source viewer shows the same list, so a
citation always points at what the officer sees.

Parsers:
- WhatsApp chat exports, Android and iPhone formats, 12/24-hour, multi-line.
- Bank alerts (SMS text, SMS backup CSV, OCR'd SMS/payment screenshots).
- Bank statements (PDF / Excel / CSV rows), direction read from Dr/Cr marks or
  from the running balance.
- Call logs (CSV with any of the common column names).
"""
import csv
import io
import re
import zipfile
from datetime import datetime, timezone

from app import storage
from app.agents import dates, extract
from app.util import IST

# ------------------------------------------------------------------ reading


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _zip_chat(data: bytes) -> str | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = [n for n in z.namelist() if n.lower().endswith(".txt") and not n.startswith("__MACOSX")]
            if not names:
                return None
            best = next((n for n in names if "chat" in n.lower()), names[0])
            return _decode(z.read(best))
    except zipfile.BadZipFile:
        return None


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [c.strftime("%d-%m-%Y %H:%M") if isinstance(c, datetime) else str(c) for c in row if c not in (None, "")]
            if cells:
                rows.append(" | ".join(cells))
    return "\n".join(rows)


def ocr_text(data: bytes) -> str:
    import pytesseract
    from PIL import Image, ImageOps
    img = Image.open(io.BytesIO(data))
    img = ImageOps.exif_transpose(img).convert("L")
    if img.width < 1000:  # phone screenshots OCR better upscaled
        img = img.resize((img.width * 2, img.height * 2))
    return pytesseract.image_to_string(img, lang="eng", config="--psm 4")


def text_lines(f: dict) -> list[str] | None:
    """Readable text of a file as lines, or None if it has no text layer."""
    if f.get("ocr") is not None:
        return f["ocr"].splitlines()
    try:
        data = storage.get(f["sha256"])
    except Exception:  # noqa: BLE001
        return None
    ext = f.get("ext", "")
    if ext in (".txt", ".csv"):
        return dates.clean(_decode(data)).splitlines()
    if ext == ".zip":
        chat = _zip_chat(data)
        return dates.clean(chat).splitlines() if chat else None
    try:
        if ext == ".pdf":
            return dates.clean(_pdf_text(data)).splitlines()
        if ext == ".xlsx":
            return dates.clean(_xlsx_text(data)).splitlines()
    except Exception:  # noqa: BLE001 - unreadable document: no events, not a crash
        return None
    return None


# -------------------------------------------------------------------- events


def ev(**kw) -> dict:
    base = {"prec": "exact", "amount": None, "account": None, "to_account": None, "upi": None, "ifsc": None,
            "phone": None, "url": None, "place": None, "person": None, "conf": 0.95, "src": [], "source_kinds": [],
            "kinds": []}
    base.update(kw)
    return base


def _dur(raw: str) -> str:
    raw = (raw or "").strip()
    parts = [int(p) for p in re.findall(r"\d+", raw)]
    if not parts:
        return ""
    if ":" in raw:
        secs = parts[-3] * 3600 + parts[-2] * 60 + parts[-1] if len(parts) >= 3 else parts[0] * 60 + parts[1]
    else:
        secs = parts[0]
    if not secs:
        return ""
    m, s = divmod(secs, 60)
    return f"{m} min {s} s" if m else f"{s} s"


# ------------------------------------------------------------------ WhatsApp

WA_HEAD = re.compile(
    r"^\[?(?P<date>\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4}),?\s+(?P<time>\d{1,2}[:.]\d{2}(?::\d{2})?(?:\s*[ap]\.?\s?m\.?)?)\]?"
    r"\s*(?:-\s*)?(?P<rest>.*)$", re.IGNORECASE)
ATTACHMENT = re.compile(r"<(media omitted|attached:[^>]*)>|^(image|video|audio|sticker|gif|document) omitted$", re.IGNORECASE)
MONEY_WORDS = re.compile(r"money (has )?gone|debited|deducted|lost|refund|withdrawn|transferred|paid", re.IGNORECASE)
OTP_WORD = re.compile(r"\b(otp|one[- ]time password|verification code)\b", re.IGNORECASE)


def parse_whatsapp(f: dict, lines: list[str], complainant: str, default_year: int) -> list[dict]:
    msgs = []  # (line_no, date_str, time_str, sender, text)
    for no, raw in enumerate(lines, start=1):
        line = raw.strip()
        m = WA_HEAD.match(line)
        if m and ": " in m.group("rest"):
            sender, text = m.group("rest").split(": ", 1)
            msgs.append([no, m.group("date"), m.group("time"), sender.strip(), text.strip()])
        elif m:
            msgs.append([no, m.group("date"), m.group("time"), None, m.group("rest")])  # system notice
        elif msgs and line:
            msgs[-1][4] += "\n" + line  # continuation of a multi-line message
    pairs = []
    for _, d, *_r in msgs:
        a, b = re.split(r"[/.\-]", d)[:2]
        pairs.append((int(a), int(b)))
    order = dates.order_for(pairs)
    senders = {s for *_x, s, _t in msgs if s}
    first = complainant.split(" ")[0].lower() if complainant else ""
    others = [s for s in senders if not extract.extract_phones(s)]
    me = {s for s in senders if (first and first in s.lower()) or s.lower() in ("you", "me")}
    if not me and len(senders) == 2 and len(others) == 1:
        me = set(others)  # a named sender opposite an unsaved number is the complainant's own phone

    events = []
    for no, d, t, sender, text in msgs:
        if not sender or ATTACHMENT.search(text):
            continue
        st = dates.stamp(f"{d} {t}", default_year, order)
        if not st:
            continue
        at, _p = st
        is_me = sender in me
        phones = extract.extract_phones(sender)
        actor = "Complainant" if is_me else (f"Caller {extract.pretty_phone(phones[0])}" if phones else sender)
        urls, amount = extract.extract_urls(text), extract.extract_amount(text)
        upis = extract.extract_upi_ids(text)
        if urls and not is_me:
            action, kinds = "Sent a link on WhatsApp", ["link"]
        elif is_me and (extract.OTP_RE.match(text) or (OTP_WORD.search(text) and re.search(r"\b\d{4,8}\b", text))):
            action, kinds = "Shared a one-time password (OTP) on WhatsApp", ["otp"]
        elif OTP_WORD.search(text) and not is_me:
            action, kinds = "Asked for a one-time password (OTP) on WhatsApp", ["otp_request"]
        elif amount or upis or MONEY_WORDS.search(text):
            short = " ".join(text.split())
            short = short if len(short) < 110 else short[:107] + "…"
            action, kinds = f"WhatsApp message: “{short}”", ["money"]
        else:
            continue
        events.append(ev(at=at, type="Message", actor=actor, action=action, amount=amount,
                         phone=phones[0] if phones else None, url=urls[0] if urls else None, upi=upis[0] if upis else None,
                         src=[{"f": f["ev_id"], "l": no}], source_kinds=["conversation"], kinds=kinds, conf=0.95))
    return events


# -------------------------------------------------------------- bank alerts

DEBIT = re.compile(r"\b(debited|debit|dr|withdrawn|withdrawal|sent|paid|transferred|trf|spent|purchase|deducted)\b", re.IGNORECASE)
CREDIT = re.compile(r"\b(credited|credit|cr|received|deposited|refund(ed)?)\b", re.IGNORECASE)
NOT_TXN = re.compile(r"\b(otp|one time password|will be debited|due on|request(ed)? money|collect request|is due|emi due|"
                     r"min(imum)? amount due|statement|offer|cashback upto|pre-?approved)\b", re.IGNORECASE)
BALANCE_NEAR = re.compile(r"(bal(ance)?|avl\.?|available)\W{0,20}$", re.IGNORECASE)
TO_VPA = re.compile(r"\b(?:to|vpa|upi[/:-])\s*(?:vpa\s*)?([\w.\-]{2,64}@[a-z]{2,32})", re.IGNORECASE)
TO_ACC = re.compile(r"\b(?:to|trf to|transfer to|benef(?:iciary)?)\s+(?:a/?c|acct?|account)?\s*(?:no\.?\s*)?([x*]{2,}\d{3,6}|\d{6,18})\b", re.IGNORECASE)
FROM_ACC = re.compile(r"\b(?:a/?c|acct?|account|card)\s*(?:no\.?\s*)?(?:ending\s*(?:with\s*)?)?([x*]{0,}\d{3,6}|[x*]{2,}\d{2,6})\b", re.IGNORECASE)
METHOD = re.compile(r"\b(IMPS|NEFT|RTGS|UPI|ATM|POS)\b", re.IGNORECASE)
REF = re.compile(r"\b(?:ref(?:erence)?|utr|txn)\s*(?:no\.?|id|#)?\s*[:.]?\s*([A-Z0-9]{6,22})\b", re.IGNORECASE)


def _txn_amount(text: str) -> float | None:
    for m in list(extract.AMOUNT_RE.finditer(text)) + list(extract.AMOUNT_WORDS_RE.finditer(text)):
        if BALANCE_NEAR.search(text[max(0, m.start() - 30):m.start()]):
            continue
        v = float(m.group(1).replace(",", ""))
        if v > 0:
            return v
    return None


def txn_from_text(text: str, default_year: int, order: str = "dmy", fallback_at: tuple | None = None) -> dict | None:
    """One bank alert or payment confirmation -> transaction fields, or None."""
    t = " ".join(dates.clean(text).split())
    if NOT_TXN.search(t):
        return None
    is_debit, is_credit = DEBIT.search(t), CREDIT.search(t)
    if not (is_debit or is_credit):
        return None
    amount = _txn_amount(t)
    if amount is None:
        return None
    st = dates.stamp(t, default_year, order) or fallback_at
    if not st:
        return None
    credit = bool(is_credit) and (not is_debit or is_credit.start() < is_debit.start())
    vpa = TO_VPA.search(t)
    upis = extract.extract_upi_ids(t)
    upi = (vpa.group(1).lower() if vpa else (upis[0] if upis else None))
    to_acc = TO_ACC.search(t)
    to_acc = to_acc.group(1).upper() if to_acc else None
    from_acc = next((m.group(1).upper() for m in FROM_ACC.finditer(t) if m.group(1).upper() != to_acc and re.search(r"\d{3}", m.group(1))), None)
    method = METHOD.search(t)
    method = method.group(1).upper() if method else ("UPI" if upi else None)
    ifsc = extract.extract_ifsc_codes(t)
    ref = REF.search(t)
    if credit:
        action = f"Credit{' by ' + method if method else ''}"
    elif upi:
        action = f"Debit by UPI to {upi}"
    elif to_acc:
        action = f"Debit by {method} to A/c {to_acc}" if method else f"Debit, transfer to A/c {to_acc}"
    else:
        action = f"Debit{' by ' + method if method else ''}"
    return {"at": st[0], "prec": st[1], "amount": amount, "account": from_acc, "to_account": None if credit else to_acc,
            "upi": None if credit else upi, "ifsc": ifsc[0] if ifsc else None, "ref": ref.group(1) if ref else None,
            "action": action, "kinds": ["credit"] if credit else ["debit"]}


def _txn_event(f: dict, line_no: int, tx: dict, source_kind: str, conf: float) -> dict:
    return ev(type="Transaction", actor=f"Complainant (A/c {tx['account']})" if tx["account"] else "Complainant",
              src=[{"f": f["ev_id"], "l": line_no}], source_kinds=[source_kind], conf=conf,
              **{k: tx[k] for k in ("at", "prec", "amount", "account", "to_account", "upi", "ifsc", "action", "kinds")})


def _order_from_text(lines: list[str]) -> str:
    pairs = [(int(m.group(1)), int(m.group(2))) for ln in lines for m in dates.NUM_DATE.finditer(ln)]
    return dates.order_for(pairs)


def parse_alert_lines(f: dict, lines: list[str], default_year: int) -> list[dict]:
    """A text file of SMS alerts, one per line (or per paragraph)."""
    order = _order_from_text(lines)
    events = []
    for no, line in enumerate(lines, start=1):
        tx = txn_from_text(line, default_year, order)
        if tx:
            events.append(_txn_event(f, no, tx, "financial_record", 0.97))
    return events


# CSV column aliases (lower-cased header -> role).
COLS = {
    "body": ("body", "message", "text", "sms", "content", "msg", "message body"),
    "date": ("date", "datetime", "date time", "timestamp", "time stamp", "received", "sent", "call date", "date/time", "readable_date"),
    "time": ("time", "call time"),
    "number": ("number", "phone", "phone number", "contact", "address", "from", "sender", "mobile", "caller", "name/number"),
    "type": ("type", "call type", "direction", "category"),
    "duration": ("duration", "duration (sec)", "duration(s)", "duration (seconds)", "call duration", "seconds"),
}


def _columns(fieldnames: list[str]) -> dict:
    low = {c.strip().lower(): c for c in fieldnames if c}
    return {role: next((low[a] for a in aliases if a in low), None) for role, aliases in COLS.items()}


def csv_rows(lines: list[str]):
    reader = csv.DictReader(io.StringIO("\n".join(lines)))
    return reader, _columns(reader.fieldnames or [])


def _cell_stamp(raw: str, default_year: int, order: str):
    raw = (raw or "").strip()
    if re.fullmatch(r"\d{10,13}", raw):  # epoch seconds / milliseconds (SMS and call backups)
        ts = int(raw) / (1000 if len(raw) == 13 else 1)
        dt = datetime.fromtimestamp(ts, tz=timezone.utc).astimezone(IST)
        return dt.isoformat(), "exact"
    return dates.stamp(raw, default_year, order)


def parse_sms_csv(f: dict, lines: list[str], default_year: int) -> list[dict]:
    reader, col = csv_rows(lines)
    order = _order_from_text(lines)
    events = []
    for no, row in enumerate(reader, start=2):
        body = (row.get(col["body"]) or "") if col["body"] else ""
        when = " ".join(x for x in ((row.get(col["date"]) or "") if col["date"] else "", (row.get(col["time"]) or "") if col["time"] else "") if x)
        tx = txn_from_text(body, default_year, order, fallback_at=_cell_stamp(when, default_year, order))
        if tx:
            events.append(_txn_event(f, no, tx, "financial_record", 0.97))
    return events


CALL_TYPE = {"1": "incoming", "2": "outgoing", "3": "missed", "5": "rejected", "6": "blocked"}


def parse_call_log(f: dict, lines: list[str], default_year: int) -> list[dict]:
    reader, col = csv_rows(lines)
    order = _order_from_text(lines)
    events = []
    for no, row in enumerate(reader, start=2):
        get = lambda role: ((row.get(col[role]) or "").strip() if col[role] else "")  # noqa: E731
        st = _cell_stamp(" ".join(x for x in (get("date"), get("time")) if x), default_year, order)
        if not st:
            continue
        number = get("number")
        ctype = CALL_TYPE.get(get("type"), get("type").lower())
        phones = extract.extract_phones(number)
        phone = phones[0] if phones else None
        who = extract.pretty_phone(phone) if phone else (number or "an unknown number")
        d_txt = _dur(get("duration"))
        if ctype.startswith("out"):
            actor, action = "Complainant", f"Called {who}" + (f" ({d_txt})" if d_txt else " (not answered)")
        elif ctype.startswith(("miss", "rej", "block")):
            actor, action = f"Caller {who}", f"{ctype.capitalize()} call to the complainant"
        else:
            actor, action = f"Caller {who}", "Called the complainant" + (f" ({d_txt})" if d_txt else "")
        events.append(ev(at=st[0], prec=st[1], type="Call", actor=actor, action=action, phone=phone, direction=ctype or "incoming",
                         src=[{"f": f["ev_id"], "l": no}], source_kinds=["call_log"], kinds=["call"], conf=0.97))
    return events


# ---------------------------------------------------------------- statements

ROW_START = re.compile(r"^\s*(\d{1,2}[/\-.](?:\d{1,2}|[A-Za-z]{3})[/\-.]\d{2,4}|\d{4}-\d{2}-\d{2}|\d{1,2}\s+[A-Za-z]{3}\s+\d{2,4})")
STMT_AMOUNT = re.compile(r"(?<![\w/.])(\d{1,3}(?:,\d{2,3})+(?:\.\d{1,2})?|\d+\.\d{2})(?![\d])")
DR = re.compile(r"\b(dr|debit|withdrawal)\b", re.IGNORECASE)
CR = re.compile(r"\b(cr|credit|deposit)\b", re.IGNORECASE)


def parse_statement(f: dict, lines: list[str], default_year: int) -> list[dict]:
    """Rows that start with a date and carry amounts. The last amount on a row is
    the running balance; the change in balance gives the direction when there
    is no Dr/Cr mark."""
    order = _order_from_text(lines)
    rows = []
    for no, line in enumerate(lines, start=1):
        if not ROW_START.match(line):
            continue
        st = dates.stamp(line, default_year, order)
        nums = [float(m.group(1).replace(",", "")) for m in STMT_AMOUNT.finditer(line)]
        if not st or not nums:
            continue
        rows.append((no, line, st, nums))
    events, prev_bal = [], None
    for no, line, st, nums in rows:
        if len(nums) == 1 and re.search(r"balance|b/f|brought forward", line, re.IGNORECASE):
            prev_bal = nums[0]  # opening / carried-forward balance row: the starting point, not a payment
            continue
        amount, bal = (nums[0], nums[-1]) if len(nums) >= 2 else (nums[0], None)
        if DR.search(line) and not CR.search(line):
            credit = False
        elif CR.search(line) and not DR.search(line):
            credit = True
        elif prev_bal is not None and bal is not None and abs(abs(bal - prev_bal) - amount) < 0.5:
            credit = bal > prev_bal
        else:  # no balance to compare (first row): fall back to the description's wording
            d_, c_ = DEBIT.search(line), CREDIT.search(line)
            credit = None if bool(d_) == bool(c_) else bool(c_)
        prev_bal = bal if bal is not None else prev_bal
        if credit is None:
            continue  # direction unknown: better no event than a wrong one
        desc = re.sub(r"\s+", " ", STMT_AMOUNT.sub("", ROW_START.sub("", line))).strip(" |")
        tx = txn_from_text(f"{'debited' if not credit else 'credited'} Rs {amount} {desc} {line[:12]}", default_year, order,
                           fallback_at=st) or {}
        events.append(ev(at=st[0], prec="day" if st[1] == "day" else st[1], type="Transaction", actor="Complainant",
                         action=(f"Credit: {desc[:60]}" if credit else (tx.get("action") or "Debit") + (f" ({desc[:50]})" if not tx.get("upi") and not tx.get("to_account") and desc else "")),
                         amount=amount, to_account=None if credit else tx.get("to_account"),
                         upi=None if credit else tx.get("upi"), kinds=["credit" if credit else "debit"],
                         src=[{"f": f["ev_id"], "l": no}], source_kinds=["financial_record"], conf=0.95))
    return events


# -------------------------------------------------------------------- images


def parse_ocr(f: dict, lines: list[str], default_year: int) -> list[dict]:
    """SMS or payment-app screenshots. Each paragraph is tried as one alert; if
    none is a complete transaction, the whole screen is tried as one receipt."""
    order = _order_from_text(lines)
    blocks, cur, start = [], [], 1
    for no, line in enumerate(lines, start=1):
        if line.strip():
            if not cur:
                start = no
            cur.append(line.strip())
        elif cur:
            blocks.append((start, " ".join(cur)))
            cur = []
    if cur:
        blocks.append((start, " ".join(cur)))
    events = []
    for start, text in blocks:
        tx = txn_from_text(text, default_year, order)
        if tx:
            events.append(_txn_event(f, start, tx, "screenshot", 0.8))
    if not events:
        whole = " ".join(t for _s, t in blocks)
        tx = txn_from_text(whole, default_year, order)
        if tx:
            amt_line = next((no for no, ln in enumerate(lines, start=1) if extract.extract_amounts(ln)), 1)
            events.append(_txn_event(f, amt_line, tx, "screenshot", 0.75))
    return events


def looks_like_statement(lines: list[str]) -> bool:
    return sum(1 for ln in lines if ROW_START.match(ln) and len(STMT_AMOUNT.findall(ln)) >= 2) >= 2


# ------------------------------------------------------------- IP trace reports

IPV4 = re.compile(r"(?<![\d.])((?:25[0-5]|2[0-4]\d|1?\d?\d)(?:\.(?:25[0-5]|2[0-4]\d|1?\d?\d)){3})(?![\d.])")
IPV6 = re.compile(r"(?<![0-9A-Fa-f:])((?:[0-9A-Fa-f]{1,4}:){2,7}[0-9A-Fa-f]{1,4})(?![0-9A-Fa-f:])")
PIN = re.compile(r"(?<![\d.:])([1-8]\d{5})(?![\d.:])")


def parse_ip_traces(f: dict, lines: list[str], default_year: int) -> list[dict]:
    """Rows of an IP trace report (ISP / CERT-In / intermediary reply): each row
    with an IP address and a 6-digit PIN code. Columns are read from a header
    row when the report has one (IP | Date/Time | ISP | PIN | Area); otherwise
    the fields are picked out of the row text."""
    order = _order_from_text(lines)
    header = None
    rows = []
    for no, raw in enumerate(lines, start=1):
        line = raw.strip()
        cells = [c.strip() for c in re.split(r"\s*\|\s*|\t", line)] if ("|" in line or "\t" in line) else None
        low = line.lower()
        if cells and "ip" in low and ("pin" in low or "postal" in low):
            header = [c.lower() for c in cells]
            continue
        ip = IPV4.search(line) or IPV6.search(line)
        if not ip:
            continue
        rest = line[:ip.start()] + " " + line[ip.end():]
        pin = PIN.search(re.sub(r"\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}|\d{1,2}:\d{2}(:\d{2})?", " ", rest))
        if not pin:
            continue
        col = {}
        if header and cells and len(cells) == len(header):
            for name, val in zip(header, cells):
                if "isp" in name or "provider" in name:
                    col["isp"] = val
                elif any(k in name for k in ("area", "location", "locality", "address", "city")):
                    col["area"] = val
        st = dates.stamp(rest, default_year, order)
        rows.append({"ip": ip.group(1), "pin": pin.group(1), "at": st[0] if st else None, "prec": st[1] if st else None,
                     "isp": col.get("isp"), "area": col.get("area"), "src": {"f": f["ev_id"], "l": no}})
    return rows
